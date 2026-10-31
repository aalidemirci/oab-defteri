"""python-docx'in TEK kapısı — evrakın Word çizicisi (T2; tasarım §9.1, TK17).

`shared/pdf.py`'nin eşidir: uygulama kodu python-docx'e yalnız buradan gider
(koruma testi `shared/tests/test_evrak.py`). Belgenin İÇERİĞİ burada kurulmaz;
içerik `shared/evrak.py`'deki modelden gelir, bu modül yalnız düzeni çizer.

Sabitlenen biçim kararları (tasarım §9.1-4, §9.1-8; iskelet envanteri §7.2):

- **Belge dili `tr-TR`** (`w:docDefaults` içindeki `w:lang`): Word'ün yazım
  denetimi ve heceleme Türkçe çalışır.
- **`w:caps` KULLANILMAZ.** Word ve LibreOffice büyük harf biçimini dile göre
  çevirir ama metnin kendisi küçük harf kalır; kopyalanan, aranan ya da başka
  programda açılan metin yanlış olur. Büyük harf verinin kendisinde gelir
  (`shared.text.tr_upper`).
- **Yazı tipi adı verilir, gömülmez:** "Times New Roman" 12 punto (RYUE md. 7).
  Dosya açıldığı sistemin yazı tipini kullanır; Pardus'ta eşdeğeri yerine geçer.
- **Sayfa A4**, kenar boşlukları PDF taban şablonuyla (`documents/base.html`) aynı.
- **Altlık** yalnız "OAB Defteri · <kod> · Sayfa n/m"dir (§9). Sayfa numarası
  Word alanıdır (`PAGE`, `NUMPAGES`): python-docx değerini hesaplamaz, bu yüzden
  altlık içerik eşdeğerliği karşılaştırmasına girmez (UY-18).

python-docx tembel içe aktarılır: modül içe aktarılınca lxml yüklenmesin.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

#: Belge dili (Word yazım denetimi ve heceleme).
DIL = "tr-TR"
#: Gövde yazı tipi ve boyutu (RYUE md. 7: Times New Roman 12 punto).
YAZI_TIPI = "Times New Roman"
GOVDE_PUNTO = 12
#: Konum cümlesinin boyutu (tasarım §9.2: 9 punto).
KONUM_PUNTO = 9
#: Kenar boşlukları (cm) — PDF taban şablonunun `@page` boşluklarıyla aynı.
KENAR_CM = {"top": 1.4, "right": 1.8, "bottom": 1.5, "left": 1.8}


def yeni_belge() -> DocxDocument:
    """A4, Türkçe, Times New Roman 12 punto boş belge.

    `Document()` varsayılan şablonu paket dizininden DOSYA YOLUYLA okur
    (`docx/templates/default.docx`); donmuş pakette bu yüzden spec
    `collect_data_files("docx")` zorunludur (tasarım §11 B10, UY-17).
    """
    from docx import Document
    from docx.shared import Cm, Mm, Pt

    belge = Document()
    _dil_ve_yazi_tipi(belge, Pt)
    for bolum in belge.sections:
        bolum.page_width = Mm(210)
        bolum.page_height = Mm(297)
        bolum.top_margin = Cm(KENAR_CM["top"])
        bolum.right_margin = Cm(KENAR_CM["right"])
        bolum.bottom_margin = Cm(KENAR_CM["bottom"])
        bolum.left_margin = Cm(KENAR_CM["left"])
    return belge


def _dil_ve_yazi_tipi(belge: DocxDocument, pt: Any) -> None:
    """Varsayılan çalıştırma özelliklerine dili, Normal biçemine yazı tipini yazar."""
    from docx.oxml.ns import qn
    from docx.oxml.parser import OxmlElement

    stiller = belge.styles.element
    varsayilan = stiller.find(qn("w:docDefaults"))
    if varsayilan is None:
        varsayilan = OxmlElement("w:docDefaults")
        stiller.insert(0, varsayilan)
    rpr_varsayilan = varsayilan.find(qn("w:rPrDefault"))
    if rpr_varsayilan is None:
        rpr_varsayilan = OxmlElement("w:rPrDefault")
        varsayilan.insert(0, rpr_varsayilan)
    rpr = rpr_varsayilan.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_varsayilan.append(rpr)
    dil = rpr.find(qn("w:lang"))
    if dil is None:
        dil = OxmlElement("w:lang")
        rpr.append(dil)
    dil.set(qn("w:val"), DIL)
    dil.set(qn("w:bidi"), DIL)

    normal = belge.styles["Normal"]
    normal.font.name = YAZI_TIPI
    normal.font.size = pt(GOVDE_PUNTO)
    # `font.name` yalnız ascii/hAnsi'yi yazar; Türkçe harfler hAnsi'dedir ama
    # karmaşık betik (cs) ve Doğu Asya yuvaları da aynı ada bağlanır ki Word
    # tema yazı tipine düşmesin.
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for yuva in ("w:cs", "w:eastAsia"):
        rfonts.set(qn(yuva), YAZI_TIPI)


def altlik_yaz(belge: DocxDocument, kod: str) -> None:
    """Her sayfanın altlığına "OAB Defteri · <kod> · Sayfa n/m" yazar (alanlarla)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    for bolum in belge.sections:
        paragraf = bolum.footer.paragraphs[0]
        paragraf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _calistir(paragraf, f"OAB Defteri · {kod} · Sayfa ", Pt(8))
        _alan(paragraf, "PAGE", Pt(8))
        _calistir(paragraf, "/", Pt(8))
        _alan(paragraf, "NUMPAGES", Pt(8))


def _calistir(paragraf: Any, metin: str, boyut: Any = None, *, kalin: bool = False) -> Any:
    run = paragraf.add_run(metin)
    if boyut is not None:
        run.font.size = boyut
    run.bold = kalin or None
    return run


# ---------------------------------------------------------------------------
# Düzen öğeleri — `shared.evrak.docx_ciz` bunları içerik modelinin sırasıyla çağırır
# ---------------------------------------------------------------------------


def ortali_satir(
    belge: DocxDocument, metin: str, *, kalin: bool = False, punto: int | None = None
) -> None:
    """Ortalı tek satır (antet, başlık, alt başlık)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    paragraf = belge.add_paragraph()
    paragraf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _calistir(paragraf, metin, Pt(punto) if punto else None, kalin=kalin)


def paragraf(belge: DocxDocument, metin: str, *, punto: int | None = None) -> None:
    """İki yana yaslı gövde paragrafı."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    p = belge.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _calistir(p, metin, Pt(punto) if punto else None)


def tablo(belge: DocxDocument, basliklar: Sequence[str], satirlar: Sequence[Sequence[str]]) -> None:
    """Çerçeveli veri tablosu; başlık satırı kalın."""
    from docx.enum.table import WD_TABLE_ALIGNMENT

    t = belge.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for hucre, metin in zip(t.rows[0].cells, basliklar, strict=True):
        _calistir(hucre.paragraphs[0], metin, kalin=True)
    for satir in satirlar:
        for hucre, metin in zip(t.add_row().cells, satir, strict=True):
            _calistir(hucre.paragraphs[0], metin)


def imza(belge: DocxDocument, tarih: str, sahipler: Sequence[tuple[str, str]]) -> None:
    """İmza ızgarası: sağda tarih, altında yan yana (ad, görev) hücreleri; çerçevesiz."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if tarih:
        p = belge.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _calistir(p, tarih)
    t = belge.add_table(rows=1, cols=len(sahipler))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for hucre, (ad, gorev) in zip(t.rows[0].cells, sahipler, strict=True):
        ad_p = hucre.paragraphs[0]
        ad_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _calistir(ad_p, ad, kalin=True)
        gorev_p = hucre.add_paragraph()
        gorev_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _calistir(gorev_p, gorev)


def _alan(paragraf: Any, komut: str, boyut: Any) -> None:
    """Basit Word alanı (`PAGE`, `NUMPAGES`): begin + instrText + separate + değer + end."""
    from docx.oxml.ns import qn
    from docx.oxml.parser import OxmlElement

    def _karakter(tur: str) -> Any:
        oge = OxmlElement("w:fldChar")
        oge.set(qn("w:fldCharType"), tur)
        return oge

    run = _calistir(paragraf, "", boyut)
    run._r.append(_karakter("begin"))
    komut_ogesi = OxmlElement("w:instrText")
    komut_ogesi.set(qn("xml:space"), "preserve")
    komut_ogesi.text = f" {komut} "
    run = _calistir(paragraf, "", boyut)
    run._r.append(komut_ogesi)
    run = _calistir(paragraf, "", boyut)
    run._r.append(_karakter("separate"))
    _calistir(paragraf, "1", boyut)  # alan güncellenene dek görünen değer
    run = _calistir(paragraf, "", boyut)
    run._r.append(_karakter("end"))


def kaydet(belge: DocxDocument) -> bytes:
    """Belgeyi .docx baytlarına çevirir (dosyaya yazmak çağıranındır)."""
    tampon = io.BytesIO()
    belge.save(tampon)
    return tampon.getvalue()


def ac(veri: bytes) -> DocxDocument:
    """.docx baytlarını python-docx belgesi olarak açar (geri okuma, testler)."""
    from docx import Document

    return Document(io.BytesIO(veri))


def govde_metin_akisi(belge: DocxDocument) -> list[str]:
    """Gövdenin metin akışı: paragraflar ve tablolar BELGE SIRASIYLA.

    `Document.paragraphs` tabloları atlar (UY-18); burada gövde öğeleri sırayla
    gezilir, tablo hücreleri satır satır, hücre içindeki paragraflar sırayla
    alınır. Altlık ve üstlük gövdede değildir, akışa girmez. Boş paragraflar
    (boşluk satırları) atlanır.
    """
    from docx.table import Table

    akis: list[str] = []
    for oge in belge.iter_inner_content():
        if isinstance(oge, Table):
            for satir in oge.rows:
                for hucre in satir.cells:
                    for paragraf in hucre.paragraphs:
                        if paragraf.text.strip():
                            akis.append(paragraf.text)
        elif oge.text.strip():
            akis.append(oge.text)
    return akis
